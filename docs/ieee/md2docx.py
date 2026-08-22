import re, sys
from docx import Document
from docx.shared import Pt

src, dst = sys.argv[1], sys.argv[2]
lines = open(src).read().split("\n")
# drop front matter
if lines[0] == "---":
    lines = lines[lines.index("---", 1) + 1:]

doc = Document()
doc.styles["Normal"].font.name = "Calibri"
doc.styles["Normal"].font.size = Pt(11)

def add_runs(par, text):
    # bold **x**, italic *x*, code `x`
    for tok in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)", text):
        if not tok:
            continue
        if tok.startswith("**"):
            par.add_run(tok[2:-2]).bold = True
        elif tok.startswith("`"):
            r = par.add_run(tok[1:-1]); r.font.name = "Consolas"
        elif tok.startswith("*"):
            par.add_run(tok[1:-1]).italic = True
        else:
            par.add_run(tok)

def flush(buf):
    if buf:
        add_runs(doc.add_paragraph(), " ".join(buf)); buf.clear()

buf, i = [], 0
while i < len(lines):
    ln = lines[i]
    if ln.startswith("|"):
        flush(buf)
        rows = []
        while i < len(lines) and lines[i].startswith("|"):
            cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
            if not all(re.fullmatch(r"-+", c) for c in cells):
                rows.append(cells)
            i += 1
        t = doc.add_table(rows=len(rows), cols=len(rows[0])); t.style = "Table Grid"
        from docx.shared import Inches
        lens = [max(len(r[c]) for r in rows) for c in range(len(rows[0]))]
        lens = [min(max(l, 8), 40) for l in lens]
        tot = sum(lens); t.autofit = False
        for c, l in enumerate(lens):
            t.columns[c].width = Inches(6.5 * l / tot)
            for r in range(len(rows)):
                t.cell(r, c).width = Inches(6.5 * l / tot)
        for r, row in enumerate(rows):
            for c, cell in enumerate(row):
                p = t.cell(r, c).paragraphs[0]; add_runs(p, cell)
                if r == 0:
                    for run in p.runs: run.bold = True
        doc.add_paragraph()
        continue
    m = re.match(r"^(#+) (.*)", ln)
    if m:
        flush(buf); doc.add_heading(m.group(2), level=min(len(m.group(1)), 3))
    elif ln.strip() == "---":
        flush(buf)
    elif re.match(r"^\s*(\d+\.|-) ", ln):
        flush(buf)
        style = "List Number" if re.match(r"^\s*\d+\.", ln) else "List Bullet"
        text = re.sub(r"^\s*(\d+\.|-) ", "", ln)
        i += 1
        while i < len(lines) and lines[i].startswith("   ") and lines[i].strip():
            text += " " + lines[i].strip(); i += 1
        add_runs(doc.add_paragraph(style=style), text)
        continue
    elif ln.strip() == "":
        flush(buf)
    else:
        buf.append(ln.strip())
    i += 1
flush(buf)
doc.save(dst)
print("wrote", dst)
